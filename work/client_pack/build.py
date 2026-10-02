"""Build the client pitch pack: content.py -> Client_Pitch_Pack.pdf (+ .html) and Client_Pitch_Pack.docx.
Run from this folder:  python3 build.py
Needs: playwright (Chromium), python-docx, Bebas Neue + Inter fonts in ~/.fonts."""
import asyncio, html, os, pathlib
import content as C

HERE = pathlib.Path(__file__).resolve().parent
IMG = HERE / 'images'
FONTS = pathlib.Path.home() / '.fonts'
e = html.escape


# ---------------------------------------------------------------- HTML / PDF
CSS = """
@font-face{font-family:Bebas;src:url('file://%(f)s/BebasNeue-Regular.ttf')}
@font-face{font-family:Inter;src:url('file://%(f)s/Inter.ttf');font-weight:100 900}
@page{size:A4;margin:0}
:root{--ink:#16181c;--paper:#f7f3ea;--red:#c0322a;--gold:#ffd66a;--muted:#6b6660;--line:#e4ddd0;--card:#fffdf8}
*{box-sizing:border-box}
html,body{margin:0;background:var(--paper);color:var(--ink);font:400 10.2pt/1.5 Inter,sans-serif;-webkit-print-color-adjust:exact;print-color-adjust:exact}
.page{width:210mm;height:297mm;overflow:hidden;padding:16mm 16mm 14mm;page-break-after:always;position:relative;background:var(--paper)}
.page:last-child{page-break-after:auto}
h1,h2,h3,.bebas{font-family:Bebas,Inter,sans-serif;font-weight:400;letter-spacing:.5px;margin:0}
h2{font-size:30pt;line-height:1}
h3{font-size:15pt;margin:11px 0 5px;color:var(--ink)}
p{margin:0 0 8px}
.cover{background:var(--ink);color:#fff;display:flex;flex-direction:column;padding:15mm 16mm 12mm}
.cover .kicker{font:600 9pt Inter;letter-spacing:3px;text-transform:uppercase;color:var(--gold)}
.cover h1{font-size:52pt;line-height:.92;margin:10px 0 6px}
.cover h1 span{color:var(--red)}
.cover .sub{font-size:13pt;color:#d8d3c8;margin-bottom:18px}
.collage{display:grid;grid-template-columns:repeat(5,1fr);gap:6px;margin:8px 0 18px}
.collage img{width:100%;height:58mm;object-fit:cover;border-radius:6px;display:block}
.cover .who{margin-top:auto;display:flex;justify-content:space-between;align-items:flex-end;border-top:1px solid #3a3d44;padding-top:12px}
.cover .who b{font-family:Bebas;font-size:24pt;font-weight:400;letter-spacing:1px}
.cover .who small{color:#aaa49a;font-size:9pt}
.intro p{color:#e9e5dc;font-size:9.6pt;margin-bottom:6px}
.proof{display:grid;grid-template-columns:1fr 1fr;gap:8px;margin-top:10px}
.proof div{background:#22252b;border-left:3px solid var(--gold);padding:8px 10px;border-radius:4px}
.proof b{display:block;font-family:Bebas;font-weight:400;font-size:15pt;color:var(--gold);letter-spacing:.5px}
.proof span{font-size:8.6pt;color:#cfcac0}
.hdr{display:flex;align-items:center;gap:12px;margin-bottom:4px}
.num{font-family:Bebas;font-size:40pt;color:var(--red);line-height:.9}
.tagline{color:var(--muted);font-size:10pt;margin:2px 0 12px}
.imgs{display:grid;gap:8px;margin:6px 0 4px}
.imgs.row3{grid-template-columns:repeat(3,1fr)} .imgs.row4{grid-template-columns:repeat(4,1fr)}
.imgs.row2{grid-template-columns:1fr 1fr;align-items:center}
.imgs img{width:100%;display:block;border-radius:6px;box-shadow:0 1px 0 rgba(0,0,0,.08)}
.imgs.row3 img,.imgs.row4 img{aspect-ratio:9/16;object-fit:cover}
.cap{font-size:8pt;color:var(--muted);margin:2px 0 10px;font-style:italic}
.hook{background:var(--ink);color:#fff;border-radius:8px;padding:11px 14px;margin:10px 0}
.hook .lbl{font:700 7.8pt Inter;letter-spacing:2px;text-transform:uppercase;color:var(--gold);margin-bottom:4px}
.hook p{margin:0;font-size:10.4pt;line-height:1.45}
.msg{border:1px dashed #c9bfae;border-radius:8px;padding:9px 13px;background:var(--card);font-size:9.2pt;color:#3b3833}
.msg .lbl{font:700 7.8pt Inter;letter-spacing:2px;text-transform:uppercase;color:var(--red);margin-bottom:3px}
.two{display:grid;grid-template-columns:1fr 1fr;gap:14px}
.chips{display:flex;flex-wrap:wrap;gap:4px}
.chip{background:#fff;border:1px solid var(--line);border-radius:20px;padding:1px 8px;font-size:8.3pt;white-space:nowrap}
.chip.main{background:var(--red);color:#fff;border-color:var(--red);font-weight:600}
.chip.tag{background:#fff3cc;border-color:#ecd68b}
ul{margin:0;padding-left:16px} li{margin:0 0 3px}
table{width:100%;border-collapse:collapse;font-size:9.2pt;background:var(--card);border-radius:8px;overflow:hidden}
th{background:var(--ink);color:#fff;text-align:left;font-weight:600;padding:6px 9px;font-size:8.4pt;letter-spacing:.5px;text-transform:uppercase}
td{padding:6px 9px;border-top:1px solid var(--line);vertical-align:top}
td.price{font-family:Bebas;font-size:17pt;color:var(--red);white-space:nowrap;width:62px}
td.tier{font-weight:700;width:72px}
.bulk{font-size:9pt;margin-top:6px;color:#3b3833}
.note{background:#fdecea;border-left:3px solid var(--red);padding:8px 11px;border-radius:4px;font-size:9pt;margin:10px 0}
.placeholder{border:2px dashed #c9bfae;border-radius:10px;padding:22px;text-align:center;color:var(--muted);margin:6px 0 12px;background:var(--card)}
.placeholder b{display:block;font-family:Bebas;font-weight:400;font-size:20pt;color:var(--ink);letter-spacing:.5px}
.glance td.price{font-size:14pt;width:auto}
.small{font-size:8.8pt;color:var(--muted)}
.steps{counter-reset:s;display:grid;grid-template-columns:repeat(3,1fr);gap:8px}
.steps div{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:6px 9px;font-size:8.5pt}
.steps b{display:block;font-family:Bebas;font-weight:400;font-size:15pt;color:var(--red)}
.foot{position:absolute;bottom:7mm;left:16mm;right:16mm;font-size:7.6pt;color:#9a948a;display:flex;justify-content:space-between}
""".replace('%(f)s', str(FONTS))


def img(name):
    return f'<img src="file://{IMG / name}">'


def chips(items, cls=''):
    return '<div class="chips">' + ''.join(f'<span class="chip {cls}">{e(i)}</span>' for i in items) + '</div>'


def foot(n):
    return f'<div class="foot"><span>{e(C.OWNER)} · {e(C.TITLE)}</span><span>{n}</span></div>'


def html_doc():
    pages = []
    # cover
    collage = ''.join(img(n) for n in ['stickman_short_1.jpg', 'lowpoly_short_2.jpg', 'stickman_short_2.jpg',
                                        'lowpoly_short_3.jpg', 'lowpoly_short_1.jpg'])
    proof = ''.join(f'<div><b>{e(a)}</b><span>{e(b)}</span></div>' for a, b in C.PROOF)
    intro = ''.join(f'<p>{e(p)}</p>' for p in C.INTRO)
    pages.append(f'''<section class="page cover">
<div class="kicker">{e(C.PREPARED_FOR)}</div>
<h1>Animated<br>Video <span>Services</span></h1>
<div class="sub">{e(C.SUBTITLE)}</div>
<div class="collage">{collage}</div>
<div class="intro">{intro}</div>
<div class="proof">{proof}</div>
<div class="who"><div><b>{e(C.OWNER)}</b><br><small>Animator &amp; video producer · English + Urdu</small></div><small>{e(C.DATE)}</small></div>
</section>''')

    # at a glance
    rows = ''.join(f'<tr><td><b>{n["num"]}</b></td><td><b>{e(n["name"])}</b><br><span class="small">{e(n["tagline"])}</span></td>'
                   f'<td class="price">{e(n["range"])}</td><td>{e(n["delivery"])}</td></tr>' for n in C.NICHES)
    heads = ''.join(f'<tr><td class="tier">{e(a)}</td><td>{e(b)}</td></tr>' for a, b in C.PROFILE_HEADLINES)
    steps = ''.join(f'<div><b>{i + 1}. {e(a)}</b>{e(b)}</div>' for i, (a, b) in enumerate(C.PROCESS))
    pages.append(f'''<section class="page">
<h2>At a glance</h2><p class="tagline">Four services, from fast and affordable to premium. Lead with the one that fits the client.</p>
<table class="glance"><tr><th>#</th><th>Service</th><th>Price</th><th>Delivery</th></tr>{rows}</table>
<h3>Profile headlines (copy &amp; paste)</h3>
<table>{heads}</table>
<h3>How every project works (tell clients this)</h3>
<div class="steps">{steps}</div>
<h3>Where to send leads</h3>
<p>When a client is interested, send Jahanzeb these details so he can confirm the price and date:</p>
<ul style="columns:2;column-gap:18px">{''.join(f"<li>{e(x)}</li>" for x in C.NEED_FROM_CLIENT)}</ul>
{foot(2)}</section>''')

    # niches
    for k, n in enumerate(C.NICHES):
        if n['images']:
            visual = f'<div class="imgs {n["image_layout"]}">{"".join(img(i) for i in n["images"])}</div><div class="cap">{e(n["caption"])}</div>'
        else:
            visual = f'<div class="placeholder"><b>Sample renders</b>{e(n.get("sample_note", ""))}</div>'
        pk = ''.join(f'<tr><td class="tier">{e(a)}</td><td class="price">{e(b)}</td><td>{e(c)}</td></tr>' for a, b, c in n['packages'])
        note = f'<div class="note">{e(n["note"])}</div>' if n.get('note') else ''
        p1 = f'''<section class="page">
<div class="hdr"><div class="num">{n["num"]}</div><h2>{e(n["name"])}</h2></div>
<div class="tagline">{e(n["tagline"])}</div>
{visual}
<p>{e(n["description"])}</p>
<div class="hook"><div class="lbl">Client hook: first message / gig intro</div><p>{e(n["hook_short"])}</p></div>
<div class="msg"><div class="lbl">Longer proposal / DM template</div>{e(n["hook_long"])}</div>
{note}
{foot(3 + k * 2)}</section>'''
        p2 = f'''<section class="page">
<div class="hdr"><div class="num">{n["num"]}</div><h2>{e(n["name"])}: selling kit</h2></div>
<h3>Main niche keyword</h3>{chips([n["main_keyword"]], 'main')}
<h3>Fiverr tags (5 max)</h3>{chips(n["tags5"], 'tag')}
<h3>All keywords for titles, descriptions, Upwork skills &amp; LinkedIn</h3>{chips(n["keywords"])}
<div class="two" style="margin-top:6px">
<div><h3>Ideal clients</h3><ul>{"".join(f"<li>{e(c)}</li>" for c in n["clients"])}</ul></div>
<div><h3>Gig / profile titles</h3><ul>{"".join(f"<li>Fiverr: {e(t)}</li>" for t in n["gig_titles"])}<li>Upwork: {e(n["upwork_title"])}</li></ul></div>
</div>
<h3>Packages &amp; price</h3>
<table><tr><th>Tier</th><th>Price</th><th>What's included</th></tr>{pk}</table>
<div class="bulk"><b>Bulk / retainer:</b> {e(n["bulk"])}</div>
{foot(4 + k * 2)}</section>'''
        pages += [p1, p2]

    # extras
    add = ''.join(f'<tr><td>{e(a)}</td><td><b>{e(b)}</b></td></tr>' for a, b in C.ADDONS)
    links = ''.join(f'<tr><td>{e(a)}</td><td>{e(b)}</td></tr>' for a, b in C.LINKS)
    n = 3 + len(C.NICHES) * 2
    pages.append(f'''<section class="page">
<h2>Add-ons, rules &amp; extras</h2><p class="tagline">Use these to upsell and to keep every job safe and profitable.</p>
<div class="two"><div><h3>Add-ons (all services)</h3><table>{add}</table></div>
<div><h3>Rules for taking jobs</h3><ul>{"".join(f"<li>{e(r)}</li>" for r in C.RULES)}</ul></div></div>
<h3>LinkedIn post ideas</h3><ul>{"".join(f"<li>{e(r)}</li>" for r in C.LINKEDIN_IDEAS)}</ul>
<h3>Portfolio links</h3><table>{links}</table>
<p class="small" style="margin-top:10px">All sample frames in this pack were rendered with our own production pipelines. Prices are in USD and are starting list prices. Quote higher for complex work, tight deadlines or extra characters.</p>
{foot(n)}</section>''')
    return f'<!doctype html><html><head><meta charset="utf-8"><title>Client Pitch Pack</title><style>{CSS}</style></head><body>{"".join(pages)}</body></html>'


async def to_pdf(html_path, pdf_path):
    from playwright.async_api import async_playwright
    exe = '/opt/pw-browsers/chromium-1194/chrome-linux/chrome'
    async with async_playwright() as p:
        b = await p.chromium.launch(executable_path=exe if os.path.exists(exe) else None)
        pg = await b.new_page()
        await pg.goto(f'file://{html_path}')
        await pg.wait_for_timeout(500)
        await pg.pdf(path=str(pdf_path), format='A4', print_background=True, margin={'top': '0', 'bottom': '0', 'left': '0', 'right': '0'})
        await b.close()


# ---------------------------------------------------------------- DOCX
def to_docx(path):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    RED = RGBColor(0xC0, 0x32, 0x2A)
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(1.8)
    d.styles['Normal'].font.name = 'Arial'; d.styles['Normal'].font.size = Pt(10.5)

    def h(text, lvl):
        p = d.add_heading(text, lvl)
        for r in p.runs: r.font.color.rgb = RED if lvl == 1 else RGBColor(0x16, 0x18, 0x1C)

    def bullets(items):
        for i in items: d.add_paragraph(i, style='List Bullet')

    def table(rows, header=None):
        t = d.add_table(rows=0, cols=len(rows[0])); t.style = 'Light Grid Accent 1'
        if header:
            c = t.add_row().cells
            for i, x in enumerate(header): c[i].text = x; c[i].paragraphs[0].runs[0].bold = True
        for r in rows:
            c = t.add_row().cells
            for i, x in enumerate(r): c[i].text = x
        d.add_paragraph()

    t = d.add_heading(C.TITLE, 0)
    d.add_paragraph(f'{C.SUBTITLE}\n{C.OWNER} · {C.PREPARED_FOR} · {C.DATE}')
    for p in C.INTRO: d.add_paragraph(p)
    h('Proof', 2)
    for a, b in C.PROOF:
        p = d.add_paragraph(style='List Bullet'); p.add_run(a + ': ').bold = True; p.add_run(b)

    h('At a glance', 1)
    table([[n['num'], n['name'], n['range'], n['delivery']] for n in C.NICHES], ['#', 'Service', 'Price', 'Delivery'])
    h('Profile headlines', 2); table([list(x) for x in C.PROFILE_HEADLINES])
    h('How every project works', 2)
    for i, (a, b) in enumerate(C.PROCESS):
        p = d.add_paragraph(style='List Number'); p.add_run(a + ': ').bold = True; p.add_run(b)
    h('What to collect from a lead', 2); bullets(C.NEED_FROM_CLIENT)

    for n in C.NICHES:
        d.add_page_break()
        h(f'{n["num"]}. {n["name"]}', 1)
        d.add_paragraph(n['tagline']).runs[0].italic = True
        if n['images']:
            w = {'row3': Cm(5.2), 'row4': Cm(3.9), 'row2': Cm(8.2)}[n['image_layout']]
            p = d.add_paragraph()
            for i in n['images']:
                p.add_run().add_picture(str(IMG / i), width=w); p.add_run('  ')
            d.add_paragraph(n['caption']).runs[0].italic = True
        else:
            d.add_paragraph('[Sample renders] ' + n.get('sample_note', ''))
        h('Description', 2); d.add_paragraph(n['description'])
        h('Client hook (first message / gig intro)', 2)
        p = d.add_paragraph(); r = p.add_run(n['hook_short']); r.bold = True
        h('Longer proposal / DM template', 2); d.add_paragraph(n['hook_long'])
        if n.get('note'): p = d.add_paragraph(); r = p.add_run('Note: ' + n['note']); r.font.color.rgb = RED
        h('Main niche keyword', 2); d.add_paragraph(n['main_keyword']).runs[0].bold = True
        h('Fiverr tags (5 max)', 2); d.add_paragraph(', '.join(n['tags5']))
        h('All keywords', 2); d.add_paragraph(', '.join(n['keywords']))
        h('Ideal clients', 2); bullets(n['clients'])
        h('Gig / profile titles', 2); bullets([f'Fiverr: {t}' for t in n['gig_titles']] + [f'Upwork: {n["upwork_title"]}'])
        h('Packages & price', 2); table([list(x) for x in n['packages']], ['Tier', 'Price', "What's included"])
        p = d.add_paragraph(); p.add_run('Bulk / retainer: ').bold = True; p.add_run(n['bulk'])

    d.add_page_break()
    h('Add-ons, rules & extras', 1)
    h('Add-ons', 2); table([list(x) for x in C.ADDONS], ['Add-on', 'Price'])
    h('Rules for taking jobs', 2); bullets(C.RULES)
    h('LinkedIn post ideas', 2); bullets(C.LINKEDIN_IDEAS)
    h('Portfolio links', 2); table([list(x) for x in C.LINKS])
    d.save(str(path))


if __name__ == '__main__':
    hp = HERE / 'Client_Pitch_Pack.html'
    hp.write_text(html_doc())
    asyncio.run(to_pdf(hp, HERE / 'Client_Pitch_Pack.pdf'))
    to_docx(HERE / 'Client_Pitch_Pack.docx')
    print('built', [p.name for p in HERE.glob('Client_Pitch_Pack.*')])
